# export_docx.py
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime

def export_to_docx(analysis: dict, output_path: str = "output.docx"):
    doc = Document()

    # Title
    title = doc.add_heading(analysis.get("title", "Digitized Document"), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Metadata block
    meta = doc.add_paragraph()
    meta.add_run(f"Type: {analysis.get('type', 'note')}  |  "
                 f"Language: {analysis.get('language', 'auto')}  |  "
                 f"Confidence: {analysis.get('confidence', 'n/a')}  |  "
                 f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    if meta.runs:
        meta.runs[0].font.color.rgb = RGBColor(0x88, 0x87, 0x80)
        meta.runs[0].font.size = Pt(9)

    doc.add_paragraph()  # spacer

    structured = analysis.get("structured", {})

    # Headings
    if structured.get("headings"):
        doc.add_heading("Headings", level=2)
        for h in structured["headings"]:
            doc.add_heading(h, level=3)

    # Paragraphs
    if structured.get("paragraphs"):
        doc.add_heading("Content", level=2)
        for para in structured["paragraphs"]:
            doc.add_paragraph(para)

    # Lists
    if structured.get("lists"):
        doc.add_heading("List Items", level=2)
        for item in structured["lists"]:
            p = doc.add_paragraph(item, style='List Bullet')

    # Math
    if structured.get("math"):
        doc.add_heading("Equations", level=2)
        for eq in structured["math"]:
            p = doc.add_paragraph()
            run = p.add_run(eq)
            run.font.name = "Courier New"
            run.font.size = Pt(12)

    # Tables
    if structured.get("tables"):
        doc.add_heading("Tables", level=2)
        for table_data in structured["tables"]:
            if table_data:
                cols = max(len(row) for row in table_data)
                table = doc.add_table(rows=len(table_data), cols=cols)
                table.style = 'Table Grid'
                for r, row in enumerate(table_data):
                    for c, cell_val in enumerate(row):
                        table.cell(r, c).text = str(cell_val)

    # Summary
    if analysis.get("summary"):
        doc.add_heading("Summary", level=2)
        p = doc.add_paragraph(analysis["summary"])
        if p.runs:
            p.runs[0].italic = True

    # Raw extracted text
    raw_text = analysis.get("raw_text", "").strip()
    if raw_text:
        doc.add_heading("Raw Extracted Text", level=2)
        raw_p = doc.add_paragraph(raw_text)
        if raw_p.runs:
            raw_p.runs[0].font.name = "Courier New"
            raw_p.runs[0].font.size = Pt(10)

    doc.save(output_path)
    print(f"Saved: {output_path}")
    return output_path
